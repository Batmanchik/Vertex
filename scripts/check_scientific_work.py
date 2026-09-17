"""Сверка чисел научной работы с файлами прогонов.

Работа объявляет правило: число со статусом [И] равно числу в файле прогона,
и расхождение считается дефектом работы, а не мелочью оформления. Правило,
которое никто не проверяет, — это пожелание. Здесь оно проверяется.

    python scripts/check_scientific_work.py

Скрипт читает artifacts/Vertex_scientific_work_v15.docx, достаёт из него
текст и сверяет каждое измеренное утверждение с тем, что лежит в artifacts.
Расхождение печатается и возвращает ненулевой код возврата.

Зачем это нужно именно здесь. Предыдущие версии работы разошлись с
измерениями не потому, что кто-то соврал, а потому, что текст писался раньше
прогонов и никто не сверял их заново после каждого пересчёта. Сверка вручную
по сорока числам не делается ни разу — значит, её надо делать командой.
"""

from __future__ import annotations

import html
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from apris.web import data  # noqa: E402

DOC = ROOT / "artifacts" / "Vertex_scientific_work_v15.docx"


def document_text() -> str:
    """Плоский текст работы без разметки."""
    with zipfile.ZipFile(DOC) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    xml = re.sub(r"</w:p>", "\n", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def fmt(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def thousands(value: float) -> str:
    """23619 -> «23 619»: в работе числа набраны с неразрывным пробелом."""
    return f"{int(round(value)):,}".replace(",", " ")


def main() -> int:
    if not DOC.exists():
        print(f"нет файла {DOC}; соберите: node scripts/make_scientific_work_v15.js")
        return 1

    text = document_text()
    # Неразрывные пробелы и типографский минус приводятся к обычным: иначе
    # сверка ловит оформление вместо чисел. «−0.0022» и «-0.0022» — одно
    # число, набранное разными знаками.
    flat = (
        text.replace(" ", " ")
        .replace(" ", " ")
        .replace("‑", "-")
        .replace("−", "-")
    )

    snap = data.snapshot()
    checks: list[tuple[str, str]] = []

    def expect(label: str, value: str) -> None:
        checks.append((label, value))

    # ── правила против модели ────────────────────────────────────────
    cells = {(c.scope, c.model): c.roc_auc for c in snap["matrix"].cells}
    expect("правила по счетам", fmt(cells[("account", "rules")]))
    expect("лес по счетам", fmt(cells[("account", "forest")]))
    expect("правила по группам", fmt(cells[("network_structural", "rules")]))
    expect("лес по группам", fmt(cells[("network_structural", "forest")]))

    # ── мир прогона ──────────────────────────────────────────────────
    world = snap["queue"].world
    expect("счетов в мире", thousands(world["accounts"]))
    expect("событий в мире", thousands(world["events"]))

    # ── лестница миров ───────────────────────────────────────────────
    for row in snap["worlds"]:
        if row.account_auc is not None:
            expect(f"лестница {row.key}", fmt(row.account_auc))

    # ── редкость ─────────────────────────────────────────────────────
    for row in snap["rarity"]:
        if row.roc_auc is not None:
            expect(f"редкость {row.prevalence:.1%} ROC-AUC", fmt(row.roc_auc))

    # ── точки порога ─────────────────────────────────────────────────
    for point in snap["points"]:
        expect("измеренная точка, сигналов", f"{point.alerts_per_1000:.1f}")
        expect("измеренная точка, точность", fmt(point.precision))
    for point in snap["projected"]:
        expect("перенос на 0.1 %, сигналов", f"{point.alerts_per_1000:.1f}")
        expect("перенос на 0.1 %, точность", fmt(point.precision))

    # ── уклонение ────────────────────────────────────────────────────
    for row in snap["evasion"]:
        if row.found_share is not None:
            expect(f"уклонение «{row.label}»", fmt(row.found_share))

    # ── Elliptic ─────────────────────────────────────────────────────
    el = snap["elliptic"]
    for arm in el.arms:
        if arm.pooled is not None:
            expect(f"Elliptic, {arm.key}", fmt(arm.pooled, 3))
    expect("Elliptic, запас над контролем", fmt(el.margin, 3))
    expect("Elliptic, узлов", thousands(el.dataset["nodes"]))
    expect("Elliptic, размеченных", thousands(el.dataset["labelled"]))
    expect("Elliptic, незаконных", thousands(el.dataset["illicit"]))

    # ── параметр W ───────────────────────────────────────────────────
    fw = snap["flow_weight"]
    expect("W, прирост к модели", f"{fw.model['lift']:.4f}")
    expect("W, важность признака", f"{fw.importance['w_fast']['mean']:.4f}")

    # ── ветви ────────────────────────────────────────────────────────
    for branch in snap["branches"]:
        expect(f"ветвь {branch.key}, между мирами", fmt(branch.across["pooled_roc_auc"]))
        expect(f"ветвь {branch.key}, внутри мира", fmt(branch.within["pooled_roc_auc"]))

    missing = [(label, value) for label, value in checks if value not in flat]

    # Проверка на присутствие ловит устаревшее число — то, которое пересчитали
    # в прогоне и забыли поправить в тексте. Но она не ловит ошибку в одной
    # ячейке, если то же число стоит рядом в прозе. Поэтому таблицы с
    # результатами сверяются поячеечно, по своему заголовку.
    wrong = check_tables(snap)
    pages, chars, rows = estimate_pages()

    print(f"проверено утверждений в тексте: {len(checks)}")
    print(f"проверено ячеек таблиц:         {CELLS_CHECKED[0]}")
    print(f"объём: {chars} символов, {rows} строк таблиц "
          f"→ около {pages:.1f} страниц при пределе {MAX_PAGES}")

    if pages > MAX_PAGES:
        print(f"ПРЕВЫШЕН ПРЕДЕЛ ОБЪЁМА: {pages:.1f} > {MAX_PAGES}")
        return 1

    if missing or wrong:
        if missing:
            print(f"НЕ НАЙДЕНО В ТЕКСТЕ: {len(missing)}")
            for label, value in missing:
                print(f"  {label:44s} ожидалось «{value}»")
        if wrong:
            print(f"РАСХОЖДЕНИЕ В ТАБЛИЦАХ: {len(wrong)}")
            for label, want, got in wrong:
                print(f"  {label:44s} ожидалось «{want}», в работе «{got}»")
        return 1
    print("все измеренные числа работы совпадают с файлами прогонов")
    return 0


CELLS_CHECKED = [0]

# Предел объёма задан требованиями к работе: не более двадцати страниц.
MAX_PAGES = 20

# Откалибровано по предыдущей версии работы целиком: 28 472 символа прозы,
# 19 строк таблиц, 12 заголовков, 134 абзаца — и 22 страницы по её же
# оглавлению. Подстановка этих величин в формулу ниже даёт 1741 символ на
# страницу при Times New Roman 14 и полуторном интервале.
CHARS_PER_PAGE = 1741
ROWS_PER_PAGE = 46

# Заголовок и отступ между абзацами занимают высоту, но почти не содержат
# символов, поэтому считать один текст мало. Ниже — их стоимость, выраженная
# в символах прозы: строка при 14 пунктах и полуторном интервале вмещает
# около 75 знаков, отступы до и после заголовка пересчитаны в такие строки.
HEADING_COST = {1: 190, 2: 165, 3: 145}
PARAGRAPH_COST = 25   # отступ после каждого абзаца

# Принудительный разрыв страницы выбрасывает остаток текущей — в среднем
# две трети. Это и была причина, по которой первая версия счётчика показывала
# двадцать страниц там, где Word показывал двадцать шесть: одиннадцать
# разрывов стоили шесть страниц, и ни один символ их не объяснял.
# Величина откалибрована по этому расхождению.
BREAK_COST = 0.67
FREE_BREAKS = 2       # после титульного листа и оглавления — они и так короткие


def count_page_breaks() -> int:
    """Принудительные разрывы страниц в собранном файле."""
    with zipfile.ZipFile(DOC) as archive:
        return archive.read("word/document.xml").decode("utf-8").count('w:type="page"')


def estimate_pages() -> tuple[float, int, int]:
    """Оценка числа страниц без отрисовки.

    LibreOffice в среде сборки не работает, а требование к объёму жёсткое,
    поэтому объём считается расчётом. Оценка нарочно осторожная: она должна
    ловить выход за предел, а не точно предсказывать вёрстку.
    """
    import docx

    document = docx.Document(DOC)
    text = sum(len(par.text) for par in document.paragraphs)
    rows = sum(len(t.rows) for t in document.tables)
    table_text = sum(len(c.text) for t in document.tables for r in t.rows for c in r.cells)

    overhead = 0
    for par in document.paragraphs:
        style = par.style.name if par.style is not None else ""
        if style.startswith("Heading"):
            overhead += HEADING_COST.get(int(style.split()[-1]), 145)
        elif par.text.strip():
            overhead += PARAGRAPH_COST

    # Титульный лист и оглавление занимают по странице каждый; каждый
    # разрыв сверх них выбрасывает остаток страницы.
    prose = text - table_text + overhead
    extra_breaks = max(0, count_page_breaks() - FREE_BREAKS)
    pages = (prose / CHARS_PER_PAGE + rows / ROWS_PER_PAGE
             + 2 + extra_breaks * BREAK_COST)
    return pages, text, rows


def check_tables(snap: dict) -> list[tuple[str, str, str]]:
    """Поячеечная сверка таблиц результатов.

    Таблица ищется по первой ячейке заголовка, а строка внутри неё — по
    первому столбцу. Привязка по смыслу, а не по номеру: вставка новой
    таблицы в работу не должна ломать проверку.
    """
    import docx

    document = docx.Document(DOC)
    tables: dict[str, dict[str, list[str]]] = {}
    for table in document.tables:
        head = table.rows[0].cells[0].text.strip()
        rows = {r.cells[0].text.strip(): [c.text.strip() for c in r.cells]
                for r in table.rows[1:]}
        tables[head] = rows

    wrong: list[tuple[str, str, str]] = []

    def cell(head: str, row: str, column: int, want: str, label: str) -> None:
        CELLS_CHECKED[0] += 1
        got = tables.get(head, {}).get(row, [])
        actual = got[column] if column < len(got) else "—строки нет—"
        if actual != want:
            wrong.append((label, want, actual))

    cells = {(c.scope, c.model): c.roc_auc for c in snap["matrix"].cells}
    cell("Уровень анализа", "По счетам", 1, fmt(cells[("account", "rules")]),
         "таблица 3: правила по счетам")
    cell("Уровень анализа", "По счетам", 3, fmt(cells[("account", "forest")]),
         "таблица 3: лес по счетам")
    cell("Уровень анализа", "По группам, структура группы", 3,
         fmt(cells[("network_structural", "forest")]), "таблица 3: лес по группам")

    for row in snap["worlds"]:
        if row.account_auc is not None:
            cell("Мир", row.key, 2, fmt(row.account_auc), f"таблица 8: {row.key}")

    for row in snap["rarity"]:
        key = f"{row.prevalence:.1%}".replace(".", ",") if False else f"{row.prevalence * 100:.1f} %"
        if row.roc_auc is not None:
            cell("Доля мошенников", key, 1, fmt(row.roc_auc), f"таблица 9: {key}")

    for row in snap["evasion"]:
        if row.found_share is not None:
            cell("Настройка уклонения", row.label, 1, fmt(row.found_share),
                 f"таблица 11: {row.label}")

    names = {"structural": "Форма, 5 признаков",
             "structural_plus_local": "Первый набор плюс активность узла",
             "shape": "Форма, 16 признаков",
             "shape_plus_local": "Расширенный набор плюс активность узла",
             "control_shuffled_labels": "Контроль: метки перемешаны"}
    for arm in snap["elliptic"].arms:
        if arm.key in names and arm.pooled is not None:
            cell("Набор признаков", names[arm.key], 1, fmt(arm.pooled),
                 f"таблица 12: {names[arm.key]}")

    labels = {"graph": "Графовая", "sequence": "Последовательностная"}
    for branch in snap["branches"]:
        cell("Ветвь", labels[branch.key], 1, fmt(branch.across["pooled_roc_auc"]),
             f"таблица 6: {labels[branch.key]}, между мирами")
        cell("Ветвь", labels[branch.key], 2, fmt(branch.within["pooled_roc_auc"]),
             f"таблица 6: {labels[branch.key]}, внутри мира")

    return wrong


if __name__ == "__main__":
    raise SystemExit(main())
